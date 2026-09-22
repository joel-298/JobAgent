"""End-to-end orchestration against the offline mock site."""
import copy
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src import runner as runner_mod
from src import state_manager as sm
from src.browser import BrowserSession
from src.excel_manager import ExcelManager
from tests import mock_pages as M
from tests.conftest import MockSite

JOB = "https://www.linkedin.com/jobs/view/555/"
COMPANY = "https://www.linkedin.com/company/acme/"
PEOPLE = "https://www.linkedin.com/company/acme/people/"
IN = "https://www.linkedin.com/in/"
FEED = "https://www.linkedin.com/feed/"

PEOPLE_LIST = [
    ("alice", "Alice A", "Technical Recruiter"),
    ("bob", "Bob B", "Talent Acquisition Lead"),
    ("carol", "Carol C", "Software Engineer"),
    ("erin", "Erin E", "Backend Developer"),
]


def _routes(job_area=M.EASY_APPLY_AREA, states=None):
    states = states or {}
    r = {
        FEED: "<html><body><nav class='global-nav'><a href='/mynetwork/'>x</a></nav></body></html>",
        JOB: M.job_page(job_area),
        COMPANY: M.people_page(PEOPLE_LIST),
        PEOPLE: M.people_page(PEOPLE_LIST),
    }
    for slug, name, headline in PEOPLE_LIST:
        r[IN + slug] = M.profile_page(name, headline, states.get(slug, "connectable"))
    return r


class MockBrowserSession(BrowserSession):
    routes: dict = {}

    def start(self):
        page = super().start()
        site = MockSite(page)
        for k, v in self.routes.items():
            site.add(k, v)
        return page


@pytest.fixture
def env(settings, tmp_path: Path, monkeypatch):
    s = copy.deepcopy(settings)
    s.paths.excel = tmp_path / "jobs.xlsx"
    s.paths.db = tmp_path / "agent.db"
    s.paths.screenshots = tmp_path / "shots"
    s.paths.logs = tmp_path / "logs"
    s.networking.hr_min, s.networking.hr_max = 1, 2
    s.networking.developer_min, s.networking.developer_max = 1, 2
    s.delays.between_companies_min = s.delays.between_companies_max = 0
    s.application.answers = {"years of experience": "4"}
    s.run.resume_paused = "no"
    monkeypatch.setattr(runner_mod, "BrowserSession", MockBrowserSession)
    monkeypatch.setattr(runner_mod, "request_intervention", lambda *a, **k: None)
    ExcelManager.create_template(s.paths.excel, sample_row=False)
    return s


def _add_row(path: Path, resume="Java Developer Resume"):
    wb = load_workbook(path)
    wb.active.append([JOB, resume, COMPANY, "HR note here", "Dev note here"])  # hr/developer_message_java
    wb.save(path)


def test_full_row_completes(env):
    MockBrowserSession.routes = _routes()
    _add_row(env.paths.excel)
    r = runner_mod.AgentRunner(env)
    assert r.run() == 1
    row = ExcelManager(env.paths.excel).load_rows()[0]
    assert row.status == "COMPLETED"
    assert row.application_status == sm.APP_COMPLETED
    assert row.hr_connections_sent == 2
    assert row.developer_connections_sent == 2
    assert row.completed_at
    run = r.state.get_run(2)
    assert run.status == "COMPLETED" and run.current_step == sm.STEP_DONE
    sent = [c["profile_url"] for c in r.state.list_connections(2) if c["result"] == sm.CONN_SENT]
    assert sorted(sent) == sorted(IN + s for s, _, _ in PEOPLE_LIST)
    r.state.close()


def test_missing_resume_pauses_then_resumes(env):
    MockBrowserSession.routes = _routes()
    _add_row(env.paths.excel, resume="Java Nope Resume")
    r = runner_mod.AgentRunner(env)
    assert r.run() == 0
    row = ExcelManager(env.paths.excel).load_rows()[0]
    assert row.status == "PAUSED"
    assert "[RESUME_SELECTION]" in row.last_error
    assert r.state.get_run(2).current_step == sm.STEP_APPLICATION
    assert r.state.list_connections(2) == []
    r.state.close()

    # user fixes the resume name in Excel and resumes
    wb = load_workbook(env.paths.excel)
    wb.active.cell(row=2, column=2, value="Java Developer Resume")
    wb.save(env.paths.excel)
    r2 = runner_mod.AgentRunner(env, resume=True)
    assert r2.run() == 1
    assert ExcelManager(env.paths.excel).load_rows()[0].status == "COMPLETED"
    r2.state.close()


def test_resume_skips_finished_phases(env):
    # Job page deliberately missing: if the agent re-applied it would fail.
    routes = _routes()
    del routes[JOB]
    MockBrowserSession.routes = routes
    _add_row(env.paths.excel)
    ExcelManager(env.paths.excel).ensure_tracking_columns()
    ExcelManager(env.paths.excel).mark_paused(2, "[HR_SEARCH] earlier failure")
    st = sm.StateManager(env.paths.db)
    st.get_or_create_run(2, JOB, COMPANY)
    st.update_run(2, status="PAUSED", application_status=sm.APP_COMPLETED, current_step=sm.STEP_HR_SEARCH, hr_successful=1, hr_attempted=1)
    st.record_connection(2, COMPANY, IN + "alice", "HR", sm.CONN_SENT)
    st.close()

    r = runner_mod.AgentRunner(env, resume=True)
    assert r.run() == 1
    row = ExcelManager(env.paths.excel).load_rows()[0]
    assert row.status == "COMPLETED"
    assert row.hr_connections_sent == 2  # alice (earlier) + bob (now)
    assert row.developer_connections_sent == 2
    hr = [c["profile_url"] for c in r.state.list_connections(2) if c["category"] == "HR"]
    assert hr.count(IN + "alice") == 1  # never re-sent
    r.state.close()


def test_never_completed_when_min_not_met(env):
    MockBrowserSession.routes = _routes(states={"alice": "connected", "bob": "connected"})
    _add_row(env.paths.excel)
    r = runner_mod.AgentRunner(env)
    assert r.run() == 0
    row = ExcelManager(env.paths.excel).load_rows()[0]
    assert row.status == "PAUSED"
    assert "[CONNECTION]" in row.last_error
    assert row.application_status == sm.APP_COMPLETED  # application progress is kept
    assert r.state.get_run(2).current_step == sm.STEP_HR_SEARCH
    r.state.close()


def test_dry_run_touches_nothing(env):
    MockBrowserSession.routes = _routes()
    _add_row(env.paths.excel)
    r = runner_mod.AgentRunner(env, dry_run=True)
    assert r.run() == 1
    row = ExcelManager(env.paths.excel).load_rows()[0]
    assert row.status == "PENDING"
    assert not env.paths.db.exists()
    r.state.close()
