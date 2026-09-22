from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.excel_manager import (
    REQUIRED_COLUMNS,
    STATUS_COMPLETED,
    STATUS_PENDING,
    TEMPLATE_MESSAGE_COLUMNS,
    TRACKING_COLUMNS,
    ExcelManager,
)

USER_COLUMNS = REQUIRED_COLUMNS + TEMPLATE_MESSAGE_COLUMNS
from src.exceptions import ExcelValidationError


@pytest.fixture
def excel(tmp_path: Path) -> ExcelManager:
    path = tmp_path / "jobs.xlsx"
    ExcelManager.create_template(path, sample_row=True)
    return ExcelManager(path)


def test_template_has_expected_headers(excel: ExcelManager):
    ws = load_workbook(excel.path).active
    headers = [c.value for c in ws[1]]
    assert headers == USER_COLUMNS + TRACKING_COLUMNS


def test_load_rows_reads_sample(excel: ExcelManager):
    rows = excel.load_rows()
    assert len(rows) == 1
    row = rows[0]
    assert row.excel_row == 2
    assert row.status == STATUS_PENDING
    assert row.resume_name == "Java Developer Resume"
    assert row.message_key == "java"
    assert "referral" in row.developer_message
    assert row.is_valid, row.problems


def test_validation_flags_bad_urls(tmp_path: Path):
    path = tmp_path / "bad.xlsx"
    ExcelManager.create_template(path, sample_row=False)
    wb = load_workbook(path)
    ws = wb.active
    ws.append(["https://example.com/not-linkedin", "", "https://www.linkedin.com/in/person", "hi", "hi"])
    ws.append(["https://www.linkedin.com/jobs/view/1", "Python Resume", "https://www.linkedin.com/company/x", "hi", "hi"])
    ws.append(["https://www.linkedin.com/jobs/view/1", "Mern Resume", "https://www.linkedin.com/company/x", "hi", "hi", "", ""])
    wb.save(path)
    rows = ExcelManager(path).load_rows()
    assert len(rows) == 3
    problems = " ".join(rows[0].problems)
    assert "job_link" in problems
    assert "resume_name is empty" in problems
    assert "company_link" in problems
    assert "no message columns match resume_name 'Python Resume'" in " ".join(rows[1].problems)
    assert "hr_message_mern is empty" in " ".join(rows[2].problems)


def test_missing_required_column_raises(tmp_path: Path):
    from openpyxl import Workbook

    path = tmp_path / "missing.xlsx"
    wb = Workbook()
    wb.active.append(["job_link", "resume_name", "company_link"])
    wb.save(path)
    with pytest.raises(ExcelValidationError):
        ExcelManager(path).load_rows()


def test_legacy_plain_message_columns_act_as_default(tmp_path: Path):
    from openpyxl import Workbook

    path = tmp_path / "legacy.xlsx"
    wb = Workbook()
    wb.active.append(["job_link", "resume_name", "company_link", "hr_message", "developer_message"])
    wb.active.append(["https://www.linkedin.com/jobs/view/1", "Any Resume", "https://www.linkedin.com/company/x", "a", "b"])
    wb.save(path)
    row = ExcelManager(path).load_rows()[0]
    assert row.is_valid, row.problems
    assert row.message_key == "default" and row.hr_message == "a" and row.developer_message == "b"


def test_ensure_tracking_columns_adds_missing(tmp_path: Path):
    from openpyxl import Workbook

    path = tmp_path / "user.xlsx"
    wb = Workbook()
    wb.active.append(USER_COLUMNS)
    wb.active.append(["https://www.linkedin.com/jobs/view/1", "Java R", "https://www.linkedin.com/company/x", "a", "b"])
    wb.save(path)
    m = ExcelManager(path)
    m.ensure_tracking_columns()
    headers = [c.value for c in load_workbook(path).active[1]]
    assert headers == USER_COLUMNS + TRACKING_COLUMNS
    rows = m.load_rows()
    assert rows[0].status == STATUS_PENDING


def test_update_never_touches_user_columns(excel: ExcelManager):
    excel.ensure_tracking_columns()
    before = excel.load_rows()[0]
    excel.mark_completed(2, hr_connections_sent=3, developer_connections_sent=4)
    after = excel.load_rows()[0]
    assert after.status == STATUS_COMPLETED
    assert after.hr_connections_sent == 3
    assert after.developer_connections_sent == 4
    assert after.completed_at
    for col in REQUIRED_COLUMNS:
        assert getattr(before, col) == getattr(after, col)
    assert before.hr_messages == after.hr_messages
    with pytest.raises(ExcelValidationError):
        excel.update_row(2, job_link="x")
    with pytest.raises(ExcelValidationError):
        excel.update_row(2, hr_message_java="x")


def test_paused_and_failed(excel: ExcelManager):
    excel.ensure_tracking_columns()
    excel.mark_paused(2, "resume not found")
    assert excel.load_rows()[0].status == "PAUSED"
    assert excel.load_rows()[0].last_error == "resume not found"
    excel.mark_failed(2, "boom")
    assert excel.load_rows()[0].status == "FAILED"
