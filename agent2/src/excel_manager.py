"""Excel input/output.

User columns: job_link, resume_name, company_link plus message columns per
resume type - hr_message_<type> / developer_message_<type> (e.g. _java,
_mern, _multi). The pair whose <type> appears in resume_name is used for that
row. Plain hr_message / developer_message act as a default pair.

User columns are never overwritten. Tracking columns are added to the header
row if missing and updated cell-by-cell with openpyxl so the user's formatting
and other content survive.
"""
from __future__ import annotations

import logging
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.exceptions import ExcelValidationError
from src.utils import is_linkedin_url, now_iso

log = logging.getLogger("agent")

REQUIRED_COLUMNS = ["job_link", "resume_name", "company_link"]
HR_MESSAGE_PREFIX = "hr_message"
DEV_MESSAGE_PREFIX = "developer_message"
DEFAULT_MESSAGE_KEY = "default"
# Columns created by --init-excel (message types you can rename/extend freely)
TEMPLATE_MESSAGE_TYPES = ["java", "mern", "multi"]
TEMPLATE_MESSAGE_COLUMNS = [
    f"{prefix}_{t}" for t in TEMPLATE_MESSAGE_TYPES for prefix in (HR_MESSAGE_PREFIX, DEV_MESSAGE_PREFIX)
]
TRACKING_COLUMNS = [
    "status",
    "application_status",
    "hr_connections_sent",
    "developer_connections_sent",
    "last_error",
    "completed_at",
]

STATUS_PENDING = "PENDING"
STATUS_IN_PROGRESS = "IN_PROGRESS"
STATUS_COMPLETED = "COMPLETED"
STATUS_PAUSED = "PAUSED"
STATUS_FAILED = "FAILED"
VALID_STATUSES = {STATUS_PENDING, STATUS_IN_PROGRESS, STATUS_COMPLETED, STATUS_PAUSED, STATUS_FAILED}
MAX_NOTE_LENGTH = 300  # LinkedIn connection-note limit


@dataclass
class JobRow:
    excel_row: int  # 1-based sheet row number (header is row 1)
    job_link: str
    resume_name: str
    company_link: str
    hr_messages: dict[str, str] = field(default_factory=dict)  # message type -> text
    developer_messages: dict[str, str] = field(default_factory=dict)
    message_key: str = ""  # resolved from resume_name
    status: str = STATUS_PENDING
    application_status: str = ""
    hr_connections_sent: int = 0
    developer_connections_sent: int = 0
    last_error: str = ""
    completed_at: str = ""
    problems: list[str] = field(default_factory=list)

    @property
    def is_valid(self) -> bool:
        return not self.problems

    @property
    def hr_message(self) -> str:
        return self.hr_messages.get(self.message_key, "")

    @property
    def developer_message(self) -> str:
        return self.developer_messages.get(self.message_key, "")

    def resolve_message_key(self) -> str | None:
        """Pick the message type whose name appears in resume_name.

        'Joel_Matthew_Java_2.pdf' + columns hr_message_java / _mern -> 'java'.
        Returns a problem string if nothing (or more than one) matches.
        """
        name = (self.resume_name or "").lower()
        keys = sorted((set(self.hr_messages) | set(self.developer_messages)) - {DEFAULT_MESSAGE_KEY})
        matches = [k for k in keys if k in name]
        if len(matches) == 1:
            self.message_key = matches[0]
            return None
        if len(matches) > 1:
            return f"resume_name matches several message types {matches} - make column names distinct"
        if DEFAULT_MESSAGE_KEY in self.hr_messages or DEFAULT_MESSAGE_KEY in self.developer_messages:
            self.message_key = DEFAULT_MESSAGE_KEY
            return None
        return f"no message columns match resume_name '{self.resume_name}' (types available: {keys or 'none'})"

    def validate(self) -> list[str]:
        p: list[str] = []
        for col in REQUIRED_COLUMNS:
            if not str(getattr(self, col) or "").strip():
                p.append(f"{col} is empty")
        problem = self.resolve_message_key() if self.resume_name else "resume_name is empty"
        if problem and self.resume_name:
            p.append(problem)
        elif not problem:
            k = self.message_key
            suffix = "" if k == DEFAULT_MESSAGE_KEY else f"_{k}"
            if not self.hr_message.strip():
                p.append(f"{HR_MESSAGE_PREFIX}{suffix} is empty")
            if not self.developer_message.strip():
                p.append(f"{DEV_MESSAGE_PREFIX}{suffix} is empty")
        if self.job_link and not is_linkedin_url(self.job_link, "/jobs/"):
            p.append("job_link is not a LinkedIn job URL (expected .../jobs/view/...)")
        if self.company_link and not is_linkedin_url(self.company_link, "/company/"):
            p.append("company_link is not a LinkedIn company URL (expected .../company/...)")
        for label, text in (("hr_message", self.hr_message), ("developer_message", self.developer_message)):
            if looks_like_placeholder(text):
                p.append(f"{label} ({self.message_key}) still contains placeholder text - fix the message before running")
            if len(text or "") > MAX_NOTE_LENGTH:
                p.append(f"{label} ({self.message_key}) is longer than {MAX_NOTE_LENGTH} characters (LinkedIn note limit)")
        if self.status and self.status not in VALID_STATUSES:
            p.append(f"unknown status '{self.status}'")
        self.problems = p
        return p


PLACEHOLDER_RE = __import__("re").compile(r"<[^>]{3,}>|PASTE|TODO|XXX|\{\{.*?\}\}", __import__("re").IGNORECASE)


def looks_like_placeholder(text: str) -> bool:
    """True for '<PASTE LINK>', 'TODO', '{{link}}' style unfinished messages."""
    return bool(text) and bool(PLACEHOLDER_RE.search(text))


def _norm_header(value) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


class ExcelManager:
    def __init__(self, path: Path):
        self.path = Path(path)
        self.sheet_name: str | None = None
        self.columns: dict[str, int] = {}  # header name -> 1-based column index
        self.hr_columns: dict[str, str] = {}  # message type -> column name
        self.dev_columns: dict[str, str] = {}

    # ---------- creation ----------
    @staticmethod
    def create_template(path: Path, sample_row: bool = True) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "jobs"
        headers = REQUIRED_COLUMNS + TEMPLATE_MESSAGE_COLUMNS + TRACKING_COLUMNS
        ws.append(headers)
        if sample_row:
            hr = (
                "Hi, I recently applied for the Software Engineer role at Example. "
                "I would love to connect and learn more about the team. Resume: <drive link>"
            )
            dev = (
                "Hi, I recently applied for the Software Engineer role at Example. "
                "If my profile looks relevant, I would really appreciate a referral. Resume: <drive link>"
            )
            row = {
                "job_link": "https://www.linkedin.com/jobs/view/123456789/",
                "resume_name": "Java Developer Resume",
                "company_link": "https://www.linkedin.com/company/example/",
                "hr_message_java": hr,
                "developer_message_java": dev,
                "status": STATUS_PENDING,
            }
            ws.append([row.get(h, "") for h in headers])
        from openpyxl.utils import get_column_letter

        for i, h in enumerate(headers, start=1):
            width = 45 if h.endswith("_link") else 60 if "message" in h else 25 if h == "resume_name" else 16
            ws.column_dimensions[get_column_letter(i)].width = width
        wb.save(path)
        return path

    # ---------- loading ----------
    def _open(self):
        if not self.path.exists():
            raise ExcelValidationError(
                f"Excel file not found: {self.path}\n"
                f"Create one with:  python app.py --init-excel"
            )
        wb = load_workbook(self.path)
        ws = wb[self.sheet_name] if self.sheet_name else wb.active
        self.sheet_name = ws.title
        return wb, ws

    def _read_headers(self, ws) -> None:
        self.columns = {}
        for idx, cell in enumerate(ws[1], start=1):
            name = _norm_header(cell.value)
            if name:
                self.columns[name] = idx
        missing = [c for c in REQUIRED_COLUMNS if c not in self.columns]
        if missing:
            raise ExcelValidationError(
                f"Excel is missing required columns: {', '.join(missing)}. "
                f"Expected header: {', '.join(REQUIRED_COLUMNS + TEMPLATE_MESSAGE_COLUMNS)}"
            )
        self.hr_columns = self._message_columns(HR_MESSAGE_PREFIX)
        self.dev_columns = self._message_columns(DEV_MESSAGE_PREFIX)
        if not self.hr_columns or not self.dev_columns:
            raise ExcelValidationError(
                "Excel needs at least one hr_message_<type> and one developer_message_<type> column "
                "(e.g. hr_message_java, developer_message_java)"
            )

    def _message_columns(self, prefix: str) -> dict[str, str]:
        """{message type: column name} for columns named <prefix> or <prefix>_<type>."""
        out: dict[str, str] = {}
        for name in self.columns:
            if name == prefix:
                out[DEFAULT_MESSAGE_KEY] = name
            elif name.startswith(prefix + "_"):
                key = name[len(prefix) + 1 :].strip("_")
                if key:
                    out[key] = name
        return out

    @property
    def user_columns(self) -> set[str]:
        return set(REQUIRED_COLUMNS) | set(self.hr_columns.values()) | set(self.dev_columns.values())

    def check_writable(self) -> None:
        """Fail early with a clear message if Excel has the file locked."""
        if not self.path.exists():
            return
        try:
            with open(self.path, "r+b"):
                pass
        except PermissionError:
            raise ExcelValidationError(
                f"Cannot write to {self.path.name} - it is open in Excel (or another program). "
                f"Close the file and run again."
            )

    def ensure_tracking_columns(self) -> None:
        """Append any missing tracking columns to the header (saves the file)."""
        self.check_writable()
        wb, ws = self._open()
        self._read_headers(ws)
        changed = False
        next_col = max(self.columns.values()) + 1
        for name in TRACKING_COLUMNS:
            if name not in self.columns:
                ws.cell(row=1, column=next_col, value=name)
                self.columns[name] = next_col
                next_col += 1
                changed = True
        if changed:
            self._save(wb)
            log.info("Added tracking columns to Excel header")

    def load_rows(self) -> list[JobRow]:
        wb, ws = self._open()
        self._read_headers(ws)
        rows: list[JobRow] = []
        for r in range(2, ws.max_row + 1):

            def get(col: str) -> str:
                idx = self.columns.get(col)
                if not idx:
                    return ""
                v = ws.cell(row=r, column=idx).value
                return "" if v is None else str(v).strip()

            if not any(get(c) for c in REQUIRED_COLUMNS):
                continue  # completely blank line

            def get_int(col: str) -> int:
                try:
                    return int(float(get(col) or 0))
                except ValueError:
                    return 0

            row = JobRow(
                excel_row=r,
                job_link=get("job_link"),
                resume_name=get("resume_name"),
                company_link=get("company_link"),
                hr_messages={k: get(col) for k, col in self.hr_columns.items()},
                developer_messages={k: get(col) for k, col in self.dev_columns.items()},
                status=(get("status") or STATUS_PENDING).upper(),
                application_status=get("application_status"),
                hr_connections_sent=get_int("hr_connections_sent"),
                developer_connections_sent=get_int("developer_connections_sent"),
                last_error=get("last_error"),
                completed_at=get("completed_at"),
            )
            row.validate()
            rows.append(row)
        return rows

    # ---------- updates ----------
    def _save(self, wb) -> None:
        backup = self.path.with_suffix(self.path.suffix + ".bak")
        try:
            shutil.copyfile(self.path, backup)
        except OSError:
            pass
        try:
            wb.save(self.path)
        except PermissionError:
            raise ExcelValidationError(
                f"Cannot write to {self.path.name} - it is open in Excel. Close it; the row's progress "
                f"is kept in data/agent.db, run again with --resume."
            )

    def update_row(self, excel_row: int, **values) -> None:
        """Write tracking values for one row. Refuses to touch the user columns."""
        for key in values:
            if key not in TRACKING_COLUMNS:
                raise ExcelValidationError(f"Refusing to write non-tracking column '{key}'")
        wb, ws = self._open()
        self._read_headers(ws)
        for key, value in values.items():
            idx = self.columns.get(key)
            if not idx:
                raise ExcelValidationError(
                    f"Tracking column '{key}' missing - call ensure_tracking_columns() first"
                )
            ws.cell(row=excel_row, column=idx, value=value)
        self._save(wb)

    # ---------- appending (used by agent1) ----------
    def existing_job_ids(self) -> set[str]:
        """LinkedIn job ids already present in the sheet (any status)."""
        import re

        ids = set()
        for row in self.load_rows():
            m = re.search(r"/jobs/view/(\d+)", row.job_link)
            if m:
                ids.add(m.group(1))
        return ids

    def append_row(self, values: dict) -> int:
        """Append a new PENDING row; keys are header names. Returns the sheet row."""
        self.ensure_tracking_columns()
        wb, ws = self._open()
        self._read_headers(ws)
        unknown = [k for k in values if k not in self.columns]
        if unknown:
            raise ExcelValidationError(f"Unknown column(s) for append: {unknown}")
        row_idx = ws.max_row + 1
        for key, value in values.items():
            ws.cell(row=row_idx, column=self.columns[key], value=value)
        ws.cell(row=row_idx, column=self.columns["status"], value=STATUS_PENDING)
        self._save(wb)
        return row_idx

    def mark_in_progress(self, excel_row: int) -> None:
        self.update_row(excel_row, status=STATUS_IN_PROGRESS, last_error="")

    def mark_paused(self, excel_row: int, error: str, **extra) -> None:
        self.update_row(excel_row, status=STATUS_PAUSED, last_error=error[:500], **extra)

    def mark_failed(self, excel_row: int, error: str, **extra) -> None:
        self.update_row(excel_row, status=STATUS_FAILED, last_error=error[:500], **extra)

    def mark_completed(self, excel_row: int, **extra) -> None:
        self.update_row(
            excel_row, status=STATUS_COMPLETED, last_error="", completed_at=now_iso(), **extra
        )
